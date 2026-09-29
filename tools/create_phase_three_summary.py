"""Create the third-phase technical summary DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"G:\Catcher\项目阶段总结\第三阶段技术总结_MinIO简历文件存储.docx")
NAVY = "1F4E78"
PALE_BLUE = "EAF3F8"
LIGHT_BORDER = "D9D9D9"
CODE_FILL = "F3F5F7"


def set_font(run, size: float, bold: bool = False, color: str = "000000", name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_style(cell, fill: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    if fill:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        tc_pr.append(shading)
    margins = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), "120")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    tc_pr.append(margins)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), LIGHT_BORDER)
        borders.append(border)
    tc_pr.append(borders)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def prevent_row_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def repeat_table_header(row) -> None:
    """Keep column labels visible when a table continues onto a new page."""

    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def style_paragraph(paragraph, before: int = 0, after: int = 12, line: float = 1.4) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_body(doc: Document, text: str, lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph)
    if lead:
        set_font(paragraph.add_run(lead), 10.5, bold=True)
    set_font(paragraph.add_run(text), 10.5)


def add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    style_paragraph(paragraph, before=32, after=10, line=1.15)
    set_font(paragraph.add_run(text), 15, bold=True)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    style_paragraph(paragraph, after=6, line=1.3)
    set_font(paragraph.add_run(text), 10.5)


def add_code(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.7)
    set_cell_style(cell, CODE_FILL)
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, after=0, line=1.25)
    set_font(paragraph.add_run(text), 9.1, name="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(15)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        set_cell_style(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(paragraph, after=0, line=1.1)
        set_font(paragraph.add_run(header), 9.8, bold=True, color="FFFFFF")
    prevent_row_split(table.rows[0])
    repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            set_cell_style(cell, PALE_BLUE if row_index % 2 else None)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(paragraph, after=0, line=1.25)
            set_font(paragraph.add_run(value), 9.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(15)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    title_style = doc.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style.paragraph_format.space_after = Pt(20)
    title_properties = title_style._element.get_or_add_pPr()
    borders = title_properties.find(qn("w:pBdr"))
    if borders is not None:
        title_properties.remove(borders)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, before=130, after=30, line=1.1)
    set_font(title.add_run("秋招求职 Agent 项目 第三阶段技术总结"), 24, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(subtitle, after=150, line=1.2)
    set_font(subtitle.add_run("MinIO 简历文件存储与受控访问"), 15)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(meta, after=170, line=1.3)
    set_font(meta.add_run("汇报对象  技术负责人\n阶段日期  2026 年 9 月 24 日"), 10.5)

    add_body(
        doc,
        "本阶段已将第二阶段的简历元数据模型接入真实 MinIO 对象存储。系统现在能够接收受限格式的简历文件，将文件写入 Bucket，并在 PostgreSQL 中以事务保存简历、版本和来源元数据。下载操作仅在确认用户归属后签发短期 URL。集成测试覆盖成功上传、格式拒绝、跨用户访问拒绝，以及数据库提交失败后的文件补偿删除。",
    )

    add_heading(doc, "一 阶段结论")
    add_table(
        doc,
        ["目标", "结果", "说明"],
        [
            ["对象存储接入", "完成", "新增 MinIO 存储适配层，负责 Bucket、上传、删除和预签名链接。"],
            ["简历原文件上传", "完成", "支持 PDF、DOCX 和 10 MB 上限，生成 SHA-256 与稳定 Object Key。"],
            ["元数据一致性", "完成", "上传后以一个 PostgreSQL 提交写入 Resume、ResumeVersion、ResumeSource。"],
            ["受控下载", "完成", "按用户归属查询简历版本后，生成 5 分钟有效的 MinIO 下载 URL。"],
            ["异常补偿验证", "完成", "写库失败时删除已上传对象；自动化测试已模拟并验证该路径。"],
        ],
        [1.45, 0.8, 4.45],
    )

    add_heading(doc, "二 本阶段范围和数据链路")
    add_body(
        doc,
        "本阶段不新增业务表，也不修改首个 Alembic 迁移。第二阶段已经建立 resumes、resume_versions 和 resume_sources，本阶段在这些表与 MinIO Bucket 之间补齐真实文件生命周期。简历内容解析、文本分块、向量化和 RAG 检索仍属于后续阶段。",
    )
    add_code(
        doc,
        "HTTP multipart upload\n"
        "  -> file validation and SHA-256\n"
        "  -> MinIO bucket career-agent-files\n"
        "  -> PostgreSQL metadata transaction\n"
        "  -> Resume + ResumeVersion + ResumeSource response",
    )
    add_body(
        doc,
        "对象采用由服务端生成的路径，而不是直接使用用户文件名。例如 resumes/{user_id}/{resume_id}/v1/original.pdf。这样 Object Key 保持唯一、稳定且不受用户输入中空格、重复名称或特殊字符影响；original_filename 仅用于显示和下载命名。",
        "对象命名  ",
    )

    add_heading(doc, "三 关键实现")
    add_table(
        doc,
        ["模块", "职责", "关键点"],
        [
            ["core/config.py", "读取对象存储配置", "MinIO 地址、凭证、Bucket、URL 时效和上传上限均来自环境变量。"],
            ["object_storage.py", "封装 MinIO SDK", "按需创建 Bucket；提供上传、删除、对象存在性验证和预签名下载链接。"],
            ["candidates.py", "实现 API 编排", "完成所有权检查、校验、上传、数据库保存和补偿删除。"],
            ["candidate.py schemas", "定义 API 契约", "新增上传结果和下载链接响应模型，保持接口输出清晰。"],
        ],
        [1.55, 1.35, 3.8],
    )
    add_bullet(doc, "格式策略：仅接受扩展名与 MIME 类型均匹配的 PDF 或 DOCX；空文件、类型不符和超过上限的文件会在写入 MinIO 前被拒绝。")
    add_bullet(doc, "完整性策略：上传时流式读取文件计算字节数和 SHA-256，随后将流复位后上传，避免额外保存临时副本。")
    add_bullet(doc, "默认状态：新版本的 source_format 为 pdf 或 docx，parse_status 为 uploaded。后续解析任务可据此接续处理。")

    add_heading(doc, "四 API 契约")
    add_table(
        doc,
        ["接口", "请求", "响应与行为"],
        [
            ["POST /users/{user_id}/resumes/upload", "multipart/form-data：file，title 可选，is_primary 可选", "创建逻辑简历、首个版本和来源元数据，并返回三者关联结果。"],
            ["GET /users/{user_id}/resume-versions/{resume_version_id}/download-url", "路径中的用户与版本 ID", "确认版本归属后返回短期 MinIO 预签名 URL 和有效秒数。"],
        ],
        [2.1, 2.25, 2.35],
    )
    add_body(
        doc,
        "当前 MVP 尚未接入登录认证，因此 user_id 只是现有 API 的归属校验边界。生产环境应从认证令牌取得当前用户身份，而不是信任客户端自行提供的 user_id。",
        "权限边界  ",
    )

    add_heading(doc, "五 一致性和异常处理")
    add_body(
        doc,
        "MinIO 与 PostgreSQL 是两个独立系统，不能共用一个数据库事务。本阶段采用先上传对象、再提交元数据的顺序，并以补偿操作处理后半段失败。数据库中的三条关联记录在同一次 Session commit 中保存，因此数据库内部保持原子性。",
    )
    add_table(
        doc,
        ["故障位置", "系统行为", "结果"],
        [
            ["校验失败", "直接返回 415、422 或 413", "不上传对象，不创建数据库记录。"],
            ["MinIO 上传失败", "返回 503", "不创建数据库记录。"],
            ["数据库提交失败", "rollback 后调用 remove_object", "正常情况下不保留数据库记录和 MinIO 对象。"],
            ["补偿删除也失败", "返回失败并保留错误上下文", "可能留下孤儿对象；后续可用定时清理任务治理。"],
        ],
        [1.45, 2.1, 3.15],
    )
    add_body(
        doc,
        "此方案适合当前同步 MVP。后续当上传量和异步处理增加时，可引入上传状态、重试策略、孤儿对象扫描和事务 Outbox，以提高崩溃恢复和可观测性。",
    )

    add_heading(doc, "六 配置和本地使用")
    add_table(
        doc,
        ["环境变量", "默认值", "作用"],
        [
            ["MINIO_BUCKET", "career-agent-files", "存放项目文件的 Bucket 名称。"],
            ["MINIO_PRESIGNED_URL_TTL_SECONDS", "300", "下载 URL 的有效期，单位为秒。"],
            ["RESUME_UPLOAD_MAX_SIZE_BYTES", "10485760", "单个简历文件最大字节数，即 10 MB。"],
            ["MINIO_ENDPOINT / MINIO_SECURE", "localhost:9000 / false", "本地 MinIO API 连接位置和传输方式。"],
        ],
        [2.2, 1.45, 3.05],
    )
    add_code(
        doc,
        "conda activate career-agent\n"
        "cd G:\\Catcher\n"
        "docker compose up -d\n"
        "fastapi dev src/career_agent/api/main.py\n"
        "# Browser: http://127.0.0.1:8000/docs",
    )

    add_heading(doc, "七 验证结果")
    add_table(
        doc,
        ["检查项", "结果", "证据"],
        [
            ["成功上传", "通过", "测试确认 MinIO 对象存在，数据库返回 object_key、文件名和 64 位 SHA-256。"],
            ["受控下载", "通过", "测试获得带 X-Amz 参数的 5 分钟预签名 URL，并拒绝不匹配用户。"],
            ["错误格式", "通过", "上传 .exe 文件得到 HTTP 415，未进入文件存储链路。"],
            ["数据库失败补偿", "通过", "测试强制 commit 报错后确认对应 MinIO 对象不存在。"],
            ["质量门禁", "通过", "pytest 5 项通过；Ruff 无问题；Pyright 0 errors；git diff --check 通过。"],
        ],
        [1.5, 0.8, 4.4],
    )

    add_heading(doc, "八 已知边界与下一阶段")
    add_bullet(doc, "MVP 不对 PDF 或 DOCX 的真实内容进行深度安全扫描；上线前应加入病毒扫描、内容解析隔离和更细的上传限流。")
    add_bullet(doc, "删除数据库记录不会自动删除 MinIO 对象。未来应提供受控删除接口，并实现失败后的异步清理。")
    add_bullet(doc, "下一阶段建议实现简历文本提取、解析任务状态、文本分块与 pgvector 向量写入，为 RAG 检索奠定可追溯的数据链路。")

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
