"""Create the second-phase technical summary DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"G:\Catcher\项目阶段总结\第二阶段技术总结_核心数据模型与数据库迁移.docx")
NAVY = "1F4E78"
PALE_BLUE = "EAF3F8"
LIGHT_BORDER = "D9D9D9"
CODE_FILL = "F3F5F7"


def set_font(run, size: float, bold: bool = False, color: str = "000000", name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), LIGHT_BORDER)


def set_cell_margins(cell, top: int = 120, start: int = 130, bottom: int = 120, end: int = 130) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cannot_split = OxmlElement("w:cantSplit")
    tr_pr.append(cannot_split)


def remove_paragraph_borders(paragraph) -> None:
    """Remove Word Title-style border residue that can render as a blue rule."""

    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def style_paragraph(paragraph, before: int = 0, after: int = 110, line: float = 1.4) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_body(doc: Document, text: str, lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph)
    if lead:
        set_font(paragraph.add_run(lead), 10.5, bold=True)
    set_font(paragraph.add_run(text), 10.5)


def add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    style_paragraph(paragraph, before=185, after=90, line=1.15)
    run = paragraph.add_run(text)
    set_font(run, 15, bold=True)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    style_paragraph(paragraph, after=70)
    set_font(paragraph.add_run(text), 10.5)


def add_code(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.7)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL)
    set_cell_borders(cell)
    set_cell_margins(cell, 100, 130, 100, 130)
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, after=0, line=1.25)
    set_font(paragraph.add_run(text), 9.2, name="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(20)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(paragraph, after=0, line=1.1)
        set_font(paragraph.add_run(header), 10, bold=True, color="FFFFFF")
    prevent_row_split(table.rows[0])

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(paragraph, after=0, line=1.25)
            set_font(paragraph.add_run(value), 9.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(20)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title_style = doc.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style.paragraph_format.space_after = Pt(20)
    style_properties = title_style._element.get_or_add_pPr()
    style_borders = style_properties.find(qn("w:pBdr"))
    if style_borders is not None:
        style_properties.remove(style_borders)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, before=130, after=30, line=1.1)
    remove_paragraph_borders(title)
    set_font(title.add_run("秋招求职 Agent 项目 第二阶段技术总结"), 24, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(subtitle, after=160, line=1.2)
    set_font(subtitle.add_run("核心数据模型与数据库迁移"), 15)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(meta, after=170, line=1.3)
    set_font(meta.add_run("汇报对象  技术负责人\n阶段日期  2026 年 9 月 24 日"), 10.5)

    add_body(
        doc,
        "本阶段已将第一阶段的数据库连通性扩展为可版本化的业务数据层。项目已建立 SQLAlchemy Session 管理、Alembic 迁移体系、六张候选人和简历核心表，以及最小可用 API。迁移已在本机 Docker PostgreSQL 上执行，API 集成测试证明数据能够真实写入、读取并在测试后清理。",
    )

    add_heading(doc, "一 阶段结论")
    add_table(
        doc,
        ["目标", "结果", "说明"],
        [
            ["数据库访问基础", "完成", "建立 Engine 缓存、Session 工厂和 FastAPI 请求级 Session 依赖。"],
            ["数据库版本管理", "完成", "初始化 Alembic，首个迁移已执行至 head。"],
            ["核心数据事实层", "完成", "创建 User、CandidateProfile、Preference、Resume、ResumeVersion、ResumeSource。"],
            ["最小业务 API", "完成", "用户、画像、偏好、简历和简历版本元数据可创建与读取。"],
            ["真实数据库验证", "完成", "集成测试通过 PostgreSQL 执行真实写入与读取，而非内存模拟。"],
        ],
        [1.55, 0.8, 4.35],
    )

    add_heading(doc, "二 设计边界与实现结构")
    add_body(
        doc,
        "本阶段只建设可信的结构化数据事实层，不引入认证、文件二进制上传、简历解析或 RAG。这样可以先稳定用户数据归属和简历版本关系，再在后续阶段把 MinIO 文件、解析文本和向量索引接入这些确定的持久化对象。",
    )
    add_code(
        doc,
        "FastAPI endpoint\n"
        "  -> Pydantic request validation\n"
        "  -> SQLAlchemy Session dependency\n"
        "  -> psycopg PostgreSQL driver\n"
        "  -> PostgreSQL tables managed by Alembic",
    )
    add_body(
        doc,
        "数据库基础代码位于 infrastructure/database.py。Engine 使用 pool_pre_ping，以便在连接池复用前探测失效连接；每个 HTTP 请求通过依赖注入获得独立 Session，并在请求结束后关闭。",
        "关键设计  ",
    )

    add_heading(doc, "三 核心数据模型")
    add_table(
        doc,
        ["实体", "关系", "职责"],
        [
            ["User", "根实体", "用户数据的归属边界；邮箱有唯一索引。"],
            ["CandidateProfile", "User 1 对 1", "保存用户确认的姓名、目标岗位、城市、经验和学历信息。"],
            ["Preference", "User 1 对 1", "保存行业、薪资和岗位类型等求职偏好。"],
            ["Resume", "User 1 对 多", "代表一份逻辑简历，可设置为主简历。"],
            ["ResumeVersion", "Resume 1 对 多", "记录每次新增版本的文本、格式、解析状态和当前版本标记。"],
            ["ResumeSource", "ResumeVersion 1 对 1", "保存 MinIO 对象键、原始文件名、媒体类型和哈希等文件元数据。"],
        ],
        [1.55, 1.45, 3.7],
    )
    add_code(
        doc,
        "User\n"
        " |- CandidateProfile  1:1\n"
        " |- Preference        1:1\n"
        " `- Resume            1:N\n"
        "      `- ResumeVersion 1:N\n"
        "           `- ResumeSource 1:1",
    )
    add_body(
        doc,
        "简历与版本分离是本阶段的关键建模选择。后续上传新文件、AI 建议和定制版本都将新增 ResumeVersion，而不是覆盖历史内容，从而让匹配分析和面试记录能够关联到当时实际使用的简历版本。",
    )

    add_heading(doc, "四 Alembic 迁移体系")
    add_body(
        doc,
        "Alembic 已配置为从应用 Settings 读取 DATABASE_URL，并加载 SQLAlchemy 的 Base.metadata。首个迁移文件 20260924_0001_create_candidate_core.py 创建了六张业务表、外键、唯一约束与必要索引。数据库中的 alembic_version 已记录当前版本。",
    )
    add_table(
        doc,
        ["命令", "用途", "本阶段结果"],
        [
            ["alembic upgrade head", "将数据库升级到最新迁移", "已成功应用首个迁移。"],
            ["alembic current", "确认数据库当前版本", "返回 20260924_0001 head。"],
            ["alembic revision", "后续生成新的结构变更记录", "作为每次模型变化的标准入口。"],
        ],
        [2.1, 2.4, 2.2],
    )
    add_body(
        doc,
        "Docker 初始化 SQL 仍仅用于启用 pgvector 扩展；业务表创建和未来字段变更均由 Alembic 管理。该边界避免了依赖清空 Docker Volume 才能应用新结构的问题。",
    )

    add_heading(doc, "五 最小业务 API")
    add_table(
        doc,
        ["接口", "作用", "持久化结果"],
        [
            ["POST /users", "创建用户", "写入 users，并阻止重复邮箱。"],
            ["GET /users/{id}", "读取用户", "按 UUID 查询用户。"],
            ["PUT GET /users/{id}/profile", "维护用户画像", "创建或更新 candidate_profiles。"],
            ["PUT GET /users/{id}/preferences", "维护求职偏好", "创建或更新 preferences。"],
            ["POST GET /users/{id}/resumes", "创建和列出简历", "写入并读取 resumes。"],
            ["POST GET /resumes/{id}/versions", "追加和列出版本", "写入 resume_versions 及可选 resume_sources。"],
        ],
        [2.45, 1.65, 2.6],
    )
    add_body(
        doc,
        "该版本中的 ResumeSource 只记录对象存储元数据，不直接接收文件内容；下一阶段会先将文件上传到 MinIO，再调用此数据模型保存关联信息。",
    )

    add_heading(doc, "六 验证结果")
    add_table(
        doc,
        ["验证项", "结果", "证据"],
        [
            ["数据库迁移", "通过", "alembic upgrade head 成功执行。"],
            ["迁移版本", "通过", "alembic current 返回 20260924_0001 head。"],
            ["业务表清单", "通过", "确认 users、candidate_profiles、preferences、resumes、resume_versions、resume_sources 和 alembic_version。"],
            ["API 集成测试", "通过", "创建用户、画像、偏好、简历、简历版本并读取；测试结束后清理临时用户。"],
            ["自动化质量检查", "通过", "pytest 2 passed；Ruff 全部通过；Pyright 0 errors。"],
            ["真实 HTTP 健康检查", "通过", "短暂启动 Uvicorn 后，GET /health 返回 status ok 与 database available。"],
        ],
        [1.7, 0.85, 4.15],
    )

    add_heading(doc, "七 当前边界与下一阶段建议")
    add_bullet(doc, "尚未实现认证、授权、租户隔离和生产级密钥管理；当前所有 API 仅适用于本机开发验证。")
    add_bullet(doc, "简历版本号在高并发场景需要进一步加锁或采用更强的幂等策略；当前实现适合 MVP 单用户开发阶段。")
    add_bullet(doc, "ResumeSource 仅保存元数据，尚未实现 MinIO 上传、病毒扫描、文件类型校验与解析任务。")
    add_bullet(doc, "Preference 的行业列表使用 PostgreSQL JSONB；后续如需复杂筛选和统计，可按业务演进为独立关联表。")
    add_body(
        doc,
        "下一阶段建议聚焦简历文件闭环：实现受控文件上传至 MinIO、创建 ResumeVersion 与 ResumeSource、提取文本并记录解析状态。完成后，RAG 的切块与 Embedding 才有可靠的数据入口和版本归属。",
        "建议  ",
    )

    add_heading(doc, "八 下一阶段完成标准")
    add_table(
        doc,
        ["完成项", "验收方式"],
        [
            ["MinIO 文件上传", "上传受允许类型和大小限制的简历，并返回稳定对象键。"],
            ["文件与版本关联", "每次上传创建 ResumeVersion 和 ResumeSource，不覆盖历史记录。"],
            ["解析状态管理", "记录上传、解析中、成功和失败等状态，并保留失败原因。"],
            ["完整性验证", "API 集成测试覆盖上传元数据、对象存在性和版本读取。"],
        ],
        [2.1, 4.55],
    )

    doc.core_properties.title = "秋招求职 Agent 项目 第二阶段技术总结"
    doc.core_properties.subject = "核心数据模型与数据库迁移"
    doc.core_properties.author = "Career Agent 项目组"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
