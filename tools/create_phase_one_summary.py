"""Create the phase-one technical summary DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"G:\Catcher\项目阶段总结\第一阶段技术总结_本地开发基础设施与健康检查.docx")

NAVY = "1F4E78"
LIGHT_BLUE = "EAF3F8"
LIGHT_GRAY = "F5F7F9"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        edge_el = borders.find(tag)
        if edge_el is None:
            edge_el = OxmlElement(f"w:{edge}")
            borders.append(edge_el)
        edge_el.set(qn("w:val"), "single")
        edge_el.set(qn("w:sz"), "6")
        edge_el.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    """Keep a table row together instead of splitting it across pages."""

    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size: float, bold: bool = False, color: RGBColor = BLACK) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def style_paragraph(paragraph, space_before: float = 0, space_after: float = 7) -> None:
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.35


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph)
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, 10.8, bold=True)
    body = paragraph.add_run(text)
    set_run_font(body, 10.8)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    style_paragraph(paragraph, space_after=3)
    for run in paragraph.runs:
        set_run_font(run, 10.6)
    run = paragraph.add_run(text)
    set_run_font(run, 10.6)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, 14 if level == 1 else 11.5, bold=True)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(paragraph, space_after=0)
        run = paragraph.add_run(header)
        set_run_font(run, 9.7, bold=True, color=RGBColor(255, 255, 255))

    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        fill = LIGHT_BLUE if row_index % 2 else "FFFFFF"
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, fill)
            set_cell_border(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            style_paragraph(paragraph, space_after=0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            set_run_font(run, 9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.22)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.1
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shading)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    run.font.color.rgb = BLACK


def create_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.8)
    title_style_pr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_style_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_style_pr.remove(title_border)

    # Cover and opening conclusion.
    title = doc.add_paragraph(style="Title")
    title_pr = title._p.get_or_add_pPr()
    direct_title_border = title_pr.find(qn("w:pBdr"))
    if direct_title_border is not None:
        title_pr.remove(direct_title_border)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(82)
    title.paragraph_format.space_after = Pt(14)
    title_run = title.add_run("秋招求职 Agent 项目 第一阶段技术总结")
    set_run_font(title_run, 22, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(46)
    subtitle_run = subtitle.add_run("本地开发基础设施与健康检查")
    set_run_font(subtitle_run, 13)

    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(40)
    run = metadata.add_run("汇报对象  技术负责人\n阶段日期  2026 年 9 月 23 日")
    set_run_font(run, 10.5)

    opening = doc.add_paragraph()
    opening.paragraph_format.space_after = Pt(16)
    opening.paragraph_format.line_spacing = 1.45
    opening_run = opening.add_run(
        "本阶段已完成项目从空目录到可运行本地开发底座的建设。Docker Desktop、PostgreSQL 含 pgvector、Redis、MinIO、Python Conda 环境及 FastAPI 健康检查均已完成实际验证。当前项目已具备稳定的依赖服务、配置边界、基础质量检查与数据库连通性，为下一阶段的数据模型和数据库迁移工作提供了可复用基础。"
    )
    set_run_font(opening_run, 11.2)

    doc.add_page_break()

    add_heading(doc, "一 阶段结论", 1)
    add_table(
        doc,
        ["目标", "结果", "说明"],
        [
            ["Windows 容器开发环境", "完成", "Docker Desktop 使用 WSL 2 后端，已可下载并运行 Linux 容器。"],
            ["本地依赖服务", "完成", "PostgreSQL 含 pgvector、Redis、MinIO 已启动并完成健康验证。"],
            ["Python 隔离环境", "完成", "Conda 环境 career-agent 使用 Python 3.11.16。"],
            ["应用最小闭环", "完成", "FastAPI 健康检查已通过实际 PostgreSQL 连接返回成功状态。"],
            ["业务数据模型与迁移", "未开始", "作为下一阶段工作，优先建设用户画像与简历的数据事实层。"],
        ],
        [1.45, 0.75, 4.75],
    )

    add_heading(doc, "二 本地运行架构", 1)
    add_body(doc, "本阶段采用 Windows 本机代码运行与 Docker 容器运行基础服务的组合。该方式降低了早期调试复杂度，同时保证 PostgreSQL、Redis、MinIO 的运行形态与后续部署环境一致。")
    add_code_block(
        doc,
        "Windows\n"
        "  └─ WSL 2\n"
        "       └─ Docker Desktop\n"
        "            ├─ PostgreSQL 17 + pgvector  localhost:5432\n"
        "            ├─ Redis 7.4                localhost:6379\n"
        "            └─ MinIO                    localhost:9000 / 9001\n\n"
        "Conda environment career-agent  Python 3.11.16\n"
        "  └─ FastAPI → SQLAlchemy + psycopg → PostgreSQL",
    )
    add_body(doc, "容器服务均通过 Docker Volume 保存数据。常规停止命令不会删除开发数据，只有显式执行带 Volume 删除参数的操作才会清空本地状态。")

    add_heading(doc, "三 基础设施实现", 1)
    add_table(
        doc,
        ["组件", "版本或镜像", "职责", "访问方式"],
        [
            ["PostgreSQL", "pgvector pg17", "核心业务数据、后续全文检索与向量检索", "localhost:5432"],
            ["pgvector", "数据库扩展", "保存 RAG 文档块的 Embedding 并执行相似度检索", "PostgreSQL 内部"],
            ["Redis", "redis 7.4 alpine", "缓存、限流、短任务状态与后续队列能力", "localhost:6379"],
            ["MinIO", "Quay 固定镜像摘要", "保存简历、JD、项目资料及后续导出文件", "API 9000 控制台 9001"],
        ],
        [1.1, 1.45, 2.85, 1.55],
    )
    add_body(doc, "PostgreSQL 首次初始化时自动执行扩展脚本，确保每个新建开发数据库均具备向量检索能力。", "关键实现  ")
    add_code_block(doc, "CREATE EXTENSION IF NOT EXISTS vector;")

    add_heading(doc, "四 配置与工程规范", 1)
    add_body(doc, "已建立 Docker Compose、环境变量模板、Git 忽略规则、Python 项目元数据、README 和初始代码目录。配置将真实密码和连接串留在本地 .env 中，并将该文件排除在版本控制之外。")
    add_table(
        doc,
        ["文件或目录", "作用"],
        [
            ["docker-compose.yml", "定义 PostgreSQL、Redis、MinIO、端口映射、Volume 与健康检查。"],
            [".env.example", "提供可共享的环境变量模板，不包含真实密码。"],
            [".env", "仅本机使用的开发凭据与 DATABASE_URL，已由 .gitignore 排除。"],
            ["pyproject.toml", "记录运行依赖、开发依赖、测试、Ruff 和 Pyright 配置。"],
            ["src/career_agent", "按 API、核心配置、领域层、基础设施层划分的应用包入口。"],
            ["tests", "自动化测试目录，当前包含健康检查的单元测试。"],
        ],
        [1.65, 5.3],
    )

    add_heading(doc, "五 Python 环境与依赖治理", 1)
    add_body(doc, "项目使用名为 career-agent 的独立 Conda 环境，而非系统 Python 或其他项目环境。实际运行 Python 为 3.11.16，项目以 editable 模式安装，代码修改可立即被解释器识别。")
    add_table(
        doc,
        ["依赖类别", "已纳入项目配置", "用途"],
        [
            ["Web 与校验", "FastAPI、Pydantic、pydantic-settings", "提供 HTTP API、请求响应模型与环境变量校验。"],
            ["数据库", "SQLAlchemy、psycopg、Alembic", "提供 ORM、PostgreSQL 驱动与后续数据库版本迁移。"],
            ["外部存储", "redis、minio", "访问缓存服务与对象存储。"],
            ["开发质量", "pytest、Ruff、Pyright", "执行自动化测试、代码规范检查和静态类型检查。"],
        ],
        [1.35, 2.6, 3.0],
    )
    add_body(doc, "PowerShell 初始执行策略阻止 Conda profile 加载。通过仅作用于当前用户的 RemoteSigned 策略恢复了本机创建脚本的加载能力，Conda 环境现已可正常激活。")

    add_heading(doc, "六 FastAPI 健康检查实现", 1)
    add_body(doc, "应用层已实现配置加载、SQLAlchemy Engine 创建与数据库探测。健康检查不是单纯返回进程存活，而是执行 SELECT 1 验证 API 到 PostgreSQL 的完整连接路径。")
    add_table(
        doc,
        ["模块", "实现内容"],
        [
            ["core config", "通过 Pydantic Settings 从 .env 加载 DATABASE_URL。"],
            ["infrastructure database", "创建启用 pool_pre_ping 的 SQLAlchemy Engine，并执行数据库探测。"],
            ["api main", "提供 GET /health；数据库不可用时返回 HTTP 503，可用时返回结构化成功结果。"],
            ["unit test", "通过替身数据库探测验证健康检查成功响应。"],
        ],
        [1.8, 5.15],
    )
    add_code_block(doc, "GET /health\n\n{\n  \"status\": \"ok\",\n  \"database\": \"available\"\n}")

    add_heading(doc, "七 网络与镜像问题处理", 1)
    add_body(doc, "初始 Docker 拉取 hello-world 时发生 Docker Hub TLS 握手超时。配置 Clash 提供的本机 HTTP 和 HTTPS 代理后，Docker Desktop 可正常拉取 Redis 与 pgvector 镜像。")
    add_body(doc, "Docker Hub 上的 minio/minio 历史镜像在当前网络环境返回访问拒绝，而非 Compose 或端口配置错误。为保持可复现性，MinIO 已切换至 Quay 的同版本固定 SHA-256 摘要镜像，而未使用不稳定的 latest 标签或来源不明的镜像加速地址。")

    add_heading(doc, "八 验证结果", 1)
    add_table(
        doc,
        ["验证项", "结果", "验证方式"],
        [
            ["Docker Desktop 与 WSL 2", "通过", "hello-world 容器完成下载、创建与执行。"],
            ["PostgreSQL 容器", "健康", "Docker Compose 状态为 healthy。"],
            ["pgvector 扩展", "通过", "数据库查询返回 vector 扩展。"],
            ["Redis", "通过", "redis-cli ping 返回 PONG。"],
            ["MinIO", "通过", "健康接口返回 HTTP 200。"],
            ["Python 数据库连接", "通过", "SQLAlchemy 使用 psycopg 成功执行 PostgreSQL 探测。"],
            ["FastAPI 健康检查", "通过", "HTTP 请求返回 status ok 和 database available。"],
            ["自动化质量检查", "通过", "pytest 1 passed；Ruff 通过；Pyright 0 errors。"],
        ],
        [1.65, 0.75, 4.55],
    )

    add_heading(doc, "九 当前边界与技术债", 1)
    add_bullet(doc, "当前尚未实现用户认证、租户隔离、权限控制和生产级密钥管理。本阶段仅覆盖本机开发环境。")
    add_bullet(doc, "尚未建立 SQLAlchemy 业务实体、Alembic 迁移目录及数据库版本管理。")
    add_bullet(doc, "尚未实现 MinIO 文件上传、简历解析、RAG 切块和 Embedding 生成。")
    add_bullet(doc, "测试中存在 FastAPI TestClient 的第三方弃用警告，当前不影响功能与静态检查，后续升级依赖时需处理。")
    add_bullet(doc, "项目依赖已记录版本范围，但尚未生成锁定文件；后续应统一选择并执行 uv 或其他锁定策略。")

    add_heading(doc, "十 下一阶段建议", 1)
    add_body(doc, "下一阶段应优先建设数据事实层，而非直接接入 Agent 或 RAG。该顺序可以确保所有后续模型输出、文件证据和用户确认操作均有稳定的持久化对象和审计边界。")
    add_table(
        doc,
        ["优先级", "工作项", "交付结果"],
        [
            ["P0", "初始化 Alembic 与数据库 Session 依赖", "可版本化的数据表迁移和受控数据库访问。"],
            ["P0", "实现 User、CandidateProfile、Resume、ResumeVersion", "用户画像与简历的真实事实层。"],
            ["P1", "生成并执行首次迁移，补充集成测试", "空数据库可重复初始化，模型可写入和读取。"],
            ["P1", "实现用户画像与简历元数据 API", "为文件上传和简历解析提供稳定入口。"],
            ["P2", "接入 MinIO 文件上传与文本提取", "形成上传简历到创建解析任务的业务起点。"],
        ],
        [0.75, 2.8, 3.4],
    )

    add_body(doc, "阶段结论  当前项目已具备可靠的本地开发基础设施与可验证的 API 数据库连通性。下一阶段将以数据模型、迁移和最小业务 API 为核心，构建用户真实求职资料的可追溯持久化基础。", "")

    doc.core_properties.title = "秋招求职 Agent 项目 第一阶段技术总结"
    doc.core_properties.subject = "本地开发基础设施与健康检查"
    doc.core_properties.author = "Career Agent 项目组"
    doc.core_properties.keywords = "Career Agent, Docker, PostgreSQL, pgvector, FastAPI"
    doc.save(OUTPUT)


if __name__ == "__main__":
    create_document()
    print(OUTPUT)
