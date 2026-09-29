"""Create the fourth-phase technical summary DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"G:\Catcher\项目阶段总结\第四阶段技术总结_简历解析与RAG基础.docx")
NAVY = "1F4E78"
PALE_BLUE = "EAF3F8"
LIGHT_BORDER = "D9D9D9"
CODE_FILL = "F3F5F7"


def set_font(
    run, size: float, bold: bool = False, color: str = "000000", name: str = "Microsoft YaHei"
) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_style(cell, fill: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    if fill:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        tc_pr.append(shading)
    margins = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), "120")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    tc_pr.append(margins)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), LIGHT_BORDER)
        borders.append(border)
    tc_pr.append(borders)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def prevent_row_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def repeat_table_header(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def style_paragraph(paragraph, before: int = 0, after: int = 12, line: float = 1.4) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_body(doc: Document, text: str, lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph)
    if lead:
        set_font(paragraph.add_run(lead), 10.5, bold=True)
    set_font(paragraph.add_run(text), 10.5)


def add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    style_paragraph(paragraph, before=32, after=10, line=1.15)
    set_font(paragraph.add_run(text), 15, bold=True)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    style_paragraph(paragraph, after=6, line=1.3)
    set_font(paragraph.add_run(text), 10.5)


def add_code(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.7)
    set_cell_style(cell, CODE_FILL)
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, after=0, line=1.25)
    set_font(paragraph.add_run(text), 9.1, name="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(15)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        set_cell_style(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(paragraph, after=0, line=1.1)
        set_font(paragraph.add_run(header), 9.8, bold=True, color="FFFFFF")
    prevent_row_split(table.rows[0])
    repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            set_cell_style(cell, PALE_BLUE if row_index % 2 else None)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(paragraph, after=0, line=1.25)
            set_font(paragraph.add_run(value), 9.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(15)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    title_style = doc.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style.paragraph_format.space_after = Pt(20)
    title_properties = title_style._element.get_or_add_pPr()
    borders = title_properties.find(qn("w:pBdr"))
    if borders is not None:
        title_properties.remove(borders)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, before=130, after=30, line=1.1)
    set_font(title.add_run("秋招求职 Agent 项目 第四阶段技术总结"), 24, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(subtitle, after=150, line=1.2)
    set_font(subtitle.add_run("简历解析 文本分块 与 RAG 向量检索基础"), 15)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(meta, after=170, line=1.3)
    set_font(meta.add_run("汇报对象  技术负责人\n阶段日期  2026 年 9 月 24 日"), 10.5)
    doc.add_page_break()

    add_body(
        doc,
        "本阶段已在已有的 MinIO 文件存储和 PostgreSQL 简历元数据之上，完成简历解析与 RAG 检索基础设施。系统可在用户显式触发后读取原文件、提取真实文本、保存解析结果、生成可追溯切块、调用 OpenAI 兼容 Embedding 服务，并使用 pgvector HNSW 索引检索当前用户的当前简历片段。自动化测试不调用真实模型服务，验证了从解析到向量检索的完整本地链路。",
    )

    add_heading(doc, "一 阶段结论")
    add_table(
        doc,
        ["目标", "结果", "说明"],
        [
            ["解析状态拆分", "完成", "parse_status 与 index_status 分别表达文本解析和向量索引状态。"],
            ["结构迁移", "完成", "第二份 Alembic 迁移已执行，新增状态列和 document_chunks 表。"],
            ["PDF DOCX 文本提取", "完成", "使用 pypdf 与 python-docx 提取文本并进行仅空白归一化。"],
            ["向量与索引", "完成", "使用 VECTOR(1536) 保存向量，HNSW 余弦索引已在 PostgreSQL 中创建。"],
            ["用户范围检索", "完成", "检索始终按 user_id 过滤，并限定为当前且已索引的简历版本。"],
            ["质量验证", "完成", "pytest 6 项通过；Ruff 与 Pyright 均无错误；迁移版本为 0002 head。"],
        ],
        [1.5, 0.8, 4.4],
    )

    add_heading(doc, "二 本阶段范围和数据链路")
    add_body(
        doc,
        "范围限定为简历原文件到可检索知识片段的基础链路，不包含 JD 匹配、生成定制简历、对话 Agent 或后台异步队列。索引必须由用户显式调用接口触发，因此不会因为上传文件而自动将文本发送至外部模型平台。",
    )
    add_code(
        doc,
        "MinIO PDF or DOCX\n"
        "  -> text extraction and whitespace normalization\n"
        "  -> resume_versions.content_text\n"
        "  -> deterministic document chunks\n"
        "  -> OpenAI compatible embeddings\n"
        "  -> PostgreSQL document_chunks plus pgvector\n"
        "  -> user scoped semantic retrieval with source metadata",
    )
    add_body(
        doc,
        "解析器只提取文档中已有文字，不改写、不补全、不推断履历事实。每个切块保留 resume_version_id、chunk_index、内容哈希、来源格式和可选章节标题，使后续回答能够回指具体简历版本和片段。",
        "可信边界  ",
    )

    add_heading(doc, "三 数据模型与迁移")
    add_table(
        doc,
        ["对象", "关键字段", "作用"],
        [
            ["resume_versions", "parse_status index_status parse_error index_error", "区分解析和索引两个步骤的进行中、成功或失败状态。"],
            ["resume_versions", "parser_version indexed_at content_text", "保留解析器版本、最后索引时间和原始提取文本，便于重建与审计。"],
            ["document_chunks", "user_id resume_version_id chunk_index content", "保存可排序、可回溯且具备用户归属的文本块。"],
            ["document_chunks", "embedding embedding_model embedding_dimensions", "保存向量及其模型和维度信息，避免不同向量空间混用。"],
            ["document_chunks", "content_hash metadata", "记录文本完整性指纹与页数、源格式等辅助来源信息。"],
        ],
        [1.55, 2.45, 2.7],
    )
    add_body(
        doc,
        "迁移文件为 20260924_0002_add_resume_rag_index.py。它会确保 vector 扩展存在，在 resume_versions 添加状态列，创建 document_chunks，并创建常规用户索引、版本索引以及 HNSW 向量索引。已执行 alembic upgrade head，当前数据库版本为 20260924_0002。",
    )
    add_code(
        doc,
        "CREATE INDEX ix_document_chunks_embedding_hnsw\n"
        "ON document_chunks\n"
        "USING hnsw (embedding vector_cosine_ops);",
    )

    add_heading(doc, "四 解析 分块 与状态转换")
    add_table(
        doc,
        ["步骤", "状态变化", "失败处理"],
        [
            ["读取源对象", "uploaded -> parsing", "MinIO 读取失败，parse_status 置为 failed 并记录 parse_error。"],
            ["提取文本", "parsing -> parsed", "PDF 无可提取文本或 DOCX 损坏时不生成内容，保留错误原因。"],
            ["准备索引", "parsed -> indexing", "保存 content_text 与 parser_version 后再调用外部 Embedding。"],
            ["写入向量", "indexing -> indexed", "Embedding 失败时 index_status 置为 failed，保留内容供后续重试。"],
        ],
        [1.25, 1.35, 4.1],
    )
    add_body(
        doc,
        "分块以段落边界为优先策略，默认大小为 800 个字符；超长段落使用带 120 字符重叠的滑动窗口。该策略保留自然段上下文，同时限制单次 Embedding 请求的输入规模。章节标题仅从原文的短行中识别，不生成新的结构化事实。",
    )
    add_bullet(doc, "重新索引某一简历版本时，仅删除该版本已有的 document_chunks，再写入新切块和新向量；其他版本不受影响。")
    add_bullet(doc, "原 MinIO 文件和 content_text 不会因索引失败而删除，因此可以在修复配置或更换模型后重新建立索引。")

    add_heading(doc, "五 Embedding 提供方配置与安全边界")
    add_body(
        doc,
        "Embedding 适配层使用 OpenAI Python SDK 的兼容调用方式，但从项目配置读取 api_key、base_url、模型名和维度。因此可以接入 OpenAI 官方服务或实现 OpenAI Embeddings 协议的聚合平台，而不让 API 路由直接依赖某一供应商。",
    )
    add_table(
        doc,
        ["配置项", "示例", "说明"],
        [
            ["EMBEDDING_BASE_URL", "https://provider.example/v1", "填写到 /v1；SDK 会自动拼接 /embeddings。"],
            ["EMBEDDING_MODEL", "text-embedding-3-small", "必须使用平台实际支持的 Embedding 模型 ID。"],
            ["EMBEDDING_DIMENSIONS", "1536", "必须与模型实际输出和数据库 VECTOR 维度一致。"],
            ["EMBEDDING_API_KEY", "仅写入 .env", "服务端私密配置，不能提交到 Git 或返回给前端。"],
        ],
        [1.9, 2.0, 2.8],
    )
    add_body(
        doc,
        "调用 Embedding 时，切块文本会发送到所配置的模型服务。现阶段应避免上传未经授权的敏感信息；后续可增加脱敏、用户同意记录、服务商审查和密钥轮换。测试环境通过 FakeEmbeddings 模拟 1536 维向量，不发送真实请求，也不消耗额度。",
        "数据外发  ",
    )

    add_heading(doc, "六 检索接口与权限约束")
    add_table(
        doc,
        ["接口", "行为", "外部调用"],
        [
            ["POST /users/{user_id}/resume-versions/{resume_version_id}/index", "显式启动读取、解析、分块和向量写入。", "会调用配置的 Embedding 服务。"],
            ["POST /users/{user_id}/knowledge/search", "将查询转换为向量，返回最相关的来源片段和距离。", "会调用配置的 Embedding 服务。"],
        ],
        [2.55, 2.75, 1.4],
    )
    add_code(
        doc,
        "SELECT id, content, embedding <=> :query_embedding AS distance\n"
        "FROM document_chunks\n"
        "WHERE user_id = :current_user_id\n"
        "  AND resume_version_is_current = true\n"
        "  AND index_status = 'indexed'\n"
        "ORDER BY embedding <=> :query_embedding\n"
        "LIMIT :limit;",
    )
    add_body(
        doc,
        "实际 SQLAlchemy 查询以 DocumentChunk.embedding.cosine_distance 生成余弦距离表达式，并连接 resume_versions 过滤当前且已索引的版本。距离越小，语义相似度越高。结果返回 chunk_id、版本 ID、切块序号、章节标题、原文和模型标识，为后续答案引用提供依据。",
    )

    add_heading(doc, "七 验证结果")
    add_table(
        doc,
        ["检查项", "结果", "证据"],
        [
            ["Alembic 迁移", "通过", "alembic current 显示 20260924_0002 head。"],
            ["数据库索引", "通过", "pg_indexes 显示 ix_document_chunks_embedding_hnsw 使用 hnsw 与 vector_cosine_ops。"],
            ["解析和索引集成", "通过", "模拟 MinIO、解析器与 Embedding 服务，验证内容、切块、状态和检索结果。"],
            ["回归测试", "通过", "pytest 6 项通过，已有用户、上传和对象存储测试均保持通过。"],
            ["静态质量", "通过", "Ruff 全部通过；Pyright 输出 0 errors；git diff --check 通过。"],
        ],
        [1.5, 0.8, 4.4],
    )

    add_heading(doc, "八 已知边界与下一阶段建议")
    add_bullet(doc, "当前索引是同步接口，适合简历文件较小的 MVP。较大文件、批量重建或重试应在后续引入 Redis 队列和后台 Worker。")
    add_bullet(doc, "VECTOR(1536) 与当前计划的模型维度绑定。若更换为其他维度的模型，需要新增或迁移向量列并整体重新索引，不能混存。")
    add_bullet(doc, "HNSW 是近似最近邻索引，优先优化检索速度。后续可加入关键词混合检索、重排序和评测集，验证结果质量。")
    add_bullet(doc, "下一阶段建议进入 JD 解析与候选人匹配：把 JD 与简历的可追溯片段接入同一检索和证据链，再实现可解释的匹配评分。")

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
