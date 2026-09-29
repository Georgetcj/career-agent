"""Safe extraction and normalization for uploaded PDF and DOCX resumes."""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from pypdf import PdfReader


class DocumentParsingError(Exception):
    """Raised when a stored resume cannot be converted into usable text."""


@dataclass(frozen=True)
class ParsedDocument:
    """Text obtained from one source document without generating new content."""

    content: str
    page_count: int | None


def normalize_extracted_text(value: str) -> str:
    """Keep original wording while making whitespace suitable for storage and chunking."""

    normalized_lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines()]
    return "\n".join(line for line in normalized_lines if line)


class ResumeDocumentParser:
    """Dispatch parsing by the already validated resume source format."""

    def parse(self, source_format: str, payload: bytes) -> ParsedDocument:
        if source_format == "pdf":
            return self._parse_pdf(payload)
        if source_format == "docx":
            return self._parse_docx(payload)
        raise DocumentParsingError(f"Unsupported resume format: {source_format}.")

    def _parse_pdf(self, payload: bytes) -> ParsedDocument:
        try:
            reader = PdfReader(BytesIO(payload))
            content = normalize_extracted_text(
                "\n".join(page.extract_text() or "" for page in reader.pages)
            )
        except Exception as error:  # pypdf exposes several parser-specific exception types.
            raise DocumentParsingError("The PDF could not be parsed.") from error
        if not content:
            raise DocumentParsingError("The PDF contains no extractable text.")
        return ParsedDocument(content=content, page_count=len(reader.pages))

    def _parse_docx(self, payload: bytes) -> ParsedDocument:
        try:
            document = Document(BytesIO(payload))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.extend(cell.text for cell in row.cells)
            content = normalize_extracted_text("\n".join(paragraphs))
        except Exception as error:  # python-docx has no stable public parse exception hierarchy.
            raise DocumentParsingError("The DOCX could not be parsed.") from error
        if not content:
            raise DocumentParsingError("The DOCX contains no extractable text.")
        return ParsedDocument(content=content, page_count=None)
